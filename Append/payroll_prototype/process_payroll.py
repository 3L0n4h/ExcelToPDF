from pathlib import Path
import re
import shutil
import sys
from openpyxl import load_workbook

# All files are expected beside this Python script.
BASE_DIR = Path(__file__).resolve().parent
MASTER_FILE = BASE_DIR / "master_list.xlsx"
TEMPLATE_FILE = BASE_DIR / "template.xlsx"
ERROR_LOG = BASE_DIR / "skipped_employees.txt"

# MOCK TEMPLATE CELL MAPPING — change these when you have the real template.
CELL_EMPLOYEE_ID = "B3"
CELL_SALARY = "B4"
CELL_JOB_DESC = "B5"
CELL_EMPLOYEE_NAME = "B6"
CELL_DEDUCTION = "B7"

# MOCK SOURCE COLUMN MAPPING (1-based column numbers).
COL_EMPLOYEE_ID = 1
COL_SALARY = 2
COL_JOB_DESC_NAME = 3
COL_DEDUCTION = 4

def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()

def safe_filename(employee_id):
    # Windows filenames cannot contain / or other reserved characters.
    # Example: 123/2026 -> 123-2026.xlsx
    return re.sub(r'[<>:"/\\\\|?*]', "-", employee_id).strip(" .-") + ".xlsx"

def parse_job_and_name(value):
    """
    Expected format: Job Description - Name
    We split only on a standalone spaced hyphen (" - ").
    Any other number of standalone separators than exactly one is flagged.
    Hyphens inside a word, e.g. Kat-kat, are not separators.
    """
    text = clean_text(value)
    if not text:
        return None, None, "Empty job/name"

    # A row of hyphens/dashes is invalid.
    if re.fullmatch(r"[-–—\s]+", text):
        return None, None, "Invalid separator only"

    # Require exactly one separator surrounded by spaces.
    parts = re.split(r"\s+-\s+", text)
    if len(parts) > 2:
        return None, None, "Multiple job/name separators"
    if len(parts) != 2:
        return None, None, "Expected 'Job Description - Name'"

    job_desc, employee_name = (part.strip() for part in parts)
    if not job_desc and not employee_name:
        return None, None, "Empty job/name"
    if not job_desc:
        return None, None, "Name only"
    if not employee_name:
        return None, None, "Job description only"
    return job_desc, employee_name, None

def get_deduction(ws, row):
    # Deduction is normally on the next row. If blank, it is 0.
    next_row = row + 1
    if next_row > ws.max_row:
        return 0
    next_id = clean_text(ws.cell(next_row, COL_EMPLOYEE_ID).value)
    next_job = clean_text(ws.cell(next_row, COL_JOB_DESC_NAME).value)
    # Do not accidentally take a deduction from the next employee's row.
    if next_id or next_job:
        return 0
    value = ws.cell(next_row, COL_DEDUCTION).value
    if value is None or clean_text(value) == "":
        return 0
    try:
        return float(value)
    except (TypeError, ValueError):
        return None

def main():
    if not MASTER_FILE.exists():
        print(f"ERROR: Source file not found: {MASTER_FILE.name}")
        return
    if not TEMPLATE_FILE.exists():
        print(f"ERROR: Template file not found: {TEMPLATE_FILE.name}")
        return

    try:
        source_wb = load_workbook(MASTER_FILE, data_only=True)
        source_ws = source_wb.active
    except Exception as exc:
        print(f"ERROR: Could not open master list: {exc}")
        return

    errors = []
    created = []
    seen_ids = set()

    # Start at row 2, assuming row 1 contains headers.
    row = 2
    while row <= source_ws.max_row:
        raw_id = source_ws.cell(row, COL_EMPLOYEE_ID).value
        raw_job_name = source_ws.cell(row, COL_JOB_DESC_NAME).value

        # Rows without an employee ID are treated as detail/deduction rows.
        if raw_id is None or clean_text(raw_id) == "":
            row += 1
            continue

        employee_id = clean_text(raw_id)
        if employee_id in seen_ids:
            errors.append(f"Employee ID {employee_id} | Row {row} | Duplicate employee ID")
            row += 1
            continue
        seen_ids.add(employee_id)

        salary = source_ws.cell(row, COL_SALARY).value
        job_desc, employee_name, reason = parse_job_and_name(raw_job_name)

        if reason:
            errors.append(f"Employee ID {employee_id} | Row {row} | {reason}")
            row += 2
            continue

        deduction = get_deduction(source_ws, row)
        if deduction is None:
            errors.append(f"Employee ID {employee_id} | Row {row + 1} | Invalid deduction; employee skipped")
            row += 2
            continue

        if salary is None or clean_text(salary) == "":
            errors.append(f"Employee ID {employee_id} | Row {row} | Missing salary")
            row += 2
            continue
        output_name = safe_filename(employee_id)
        output_path = BASE_DIR / output_name
        if output_path.exists():
            errors.append(f"Employee ID {employee_id} | Row {row} | Output already exists ({output_name}); not overwritten")
            row += 2
            continue

        try:
            # Load a fresh copy of the template for each employee.
            template_wb = load_workbook(TEMPLATE_FILE)
            template_ws = template_wb.active
            template_ws[CELL_EMPLOYEE_ID] = employee_id
            template_ws[CELL_SALARY] = salary
            template_ws[CELL_JOB_DESC] = job_desc
            template_ws[CELL_EMPLOYEE_NAME] = employee_name
            template_ws[CELL_DEDUCTION] = deduction
            template_wb.save(output_path)
            template_wb.close()
            created.append(output_name)
        except Exception as exc:
            errors.append(f"Employee ID {employee_id} | Row {row} | Could not save template: {exc}")

        row += 2

    source_wb.close()

    # Always write a log, even if there were no errors.
    if errors:
        ERROR_LOG.write_text("\n".join(errors) + "\n", encoding="utf-8")
    else:
        ERROR_LOG.write_text("No skipped employees or errors.\n", encoding="utf-8")

    print(f"Finished. Created {len(created)} pay sheet(s).")
    for filename in created:
        print(f"  Created: {filename}")
    print(f"Error log: {ERROR_LOG.name}")
    if errors:
        print(f"Items requiring review: {len(errors)}")

if __name__ == "__main__":
    main()
